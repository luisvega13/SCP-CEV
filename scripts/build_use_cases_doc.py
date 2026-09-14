from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Casos_de_Uso_Portal_Escolar.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
SKY = "0EA5E9"
MUTED = "5B6B82"
LIGHT = "E8EEF5"
LIGHTER = "F4F6F9"
WHITE = "FFFFFF"
INK = "172033"
GREEN = "166534"
AMBER = "92400E"
RED = "9B1C1C"
BORDER = "CBD5E1"


USE_CASES = [
    {
        "module": "Acceso y seguridad",
        "id": "CU-01",
        "name": "Iniciar sesión",
        "actor": "Administrador o alumno",
        "support": "Supabase Auth",
        "objective": "Autenticar al usuario y dirigirlo al panel correspondiente a su rol.",
        "trigger": "El usuario captura correo y contraseña y selecciona Entrar.",
        "pre": ["La cuenta existe en Supabase Auth.", "La cuenta tiene el rol admin o student en app_metadata."],
        "post": ["La sesión queda activa mediante cookies seguras.", "El usuario llega al panel autorizado para su rol."],
        "flow": [
            "El sistema valida que el correo tenga formato correcto y que la contraseña no esté vacía.",
            "La aplicación envía las credenciales a Supabase Auth con un límite de espera de 15 segundos.",
            "Supabase valida las credenciales y devuelve el usuario autenticado.",
            "El sistema verifica el rol guardado en app_metadata.",
            "Si el rol es admin, redirige a /dashboard/admin; si es student, a /dashboard/alumno.",
            "La interfaz confirma la navegación mostrando el estado de carga durante el proceso.",
        ],
        "alternate": [
            "Credenciales incorrectas: se informa que el correo o la contraseña no coinciden.",
            "Problema de red o tiempo agotado: se muestra un mensaje de conexión y se conserva el formulario.",
            "Rol ausente o inválido: se cierra la sesión y se impide el acceso.",
        ],
        "rules": ["RN-01", "RN-02", "RN-03"],
        "evidence": "app/login/page.tsx; middleware.ts",
    },
    {
        "module": "Acceso y seguridad",
        "id": "CU-02",
        "name": "Proteger rutas por sesión y rol",
        "actor": "Sistema",
        "support": "Middleware de Next.js y Supabase Auth",
        "objective": "Evitar que una sesión no autenticada o con un rol incorrecto acceda a áreas privadas.",
        "trigger": "Se solicita una URL bajo /dashboard.",
        "pre": ["La aplicación dispone de las variables públicas de conexión a Supabase."],
        "post": ["Solo se entrega la ruta que corresponde al rol confirmado."],
        "flow": [
            "El middleware solicita los claims de la sesión a Supabase.",
            "Si no existe identidad confirmada, redirige a /login.",
            "Si un alumno solicita una ruta administrativa, lo redirige a su panel.",
            "Si un administrador solicita el panel del alumno, lo redirige al panel administrativo.",
            "Si la ruta y el rol son compatibles, permite continuar la solicitud.",
        ],
        "alternate": ["Supabase no puede confirmar la identidad: una ruta privada se rechaza y el usuario vuelve a iniciar sesión."],
        "rules": ["RN-01", "RN-02"],
        "evidence": "middleware.ts; app/dashboard/admin/layout.tsx; app/dashboard/alumno/layout.tsx",
    },
    {
        "module": "Acceso y seguridad",
        "id": "CU-03",
        "name": "Cerrar sesión",
        "actor": "Administrador o alumno",
        "support": "Supabase Auth",
        "objective": "Finalizar la sesión activa y retirar el acceso a las rutas privadas.",
        "trigger": "El usuario selecciona Cerrar sesión en el menú lateral.",
        "pre": ["Existe una sesión activa."],
        "post": ["La sesión se invalida y la aplicación muestra /login."],
        "flow": ["La interfaz marca la acción como en curso.", "Supabase cierra la sesión.", "La aplicación reemplaza la ruta por /login y actualiza el estado del navegador."],
        "alternate": ["Si el cierre remoto falla, la interfaz informa el error y permite reintentar."],
        "rules": ["RN-03"],
        "evidence": "components/Sidebar.tsx; app/cuenta-suspendida/page.tsx",
    },
    {
        "module": "Administración general",
        "id": "CU-04",
        "name": "Consultar el resumen administrativo",
        "actor": "Administrador",
        "support": "Función obtener_resumen_administrativo",
        "objective": "Ofrecer una vista ejecutiva de la situación académica, financiera y operativa del ciclo.",
        "trigger": "El administrador abre el panel principal o selecciona Actualizar.",
        "pre": ["La sesión pertenece a un administrador.", "Existe un ciclo escolar válido."],
        "post": ["Se muestran métricas consolidadas y la hora de actualización."],
        "flow": [
            "El sistema aplica la regla de pausas por inscripción vencida.",
            "Consulta el resumen del ciclo mediante una función protegida de base de datos.",
            "Presenta alumnos totales, activos, en pausa y dados de baja.",
            "Presenta recaudación del mes, del día y del ciclo, ticket promedio, cartera y eficiencia de cobranza.",
            "Desglosa cargos, niveles, tipos de cobro, métodos de pago, tendencia, actividad reciente y auditoría.",
            "Permite actualizar los datos bajo demanda y muestra estados de carga o error.",
        ],
        "alternate": ["Ciclo inválido o acceso no administrativo: la función rechaza la consulta.", "Error de red: se conserva una vista de error con opción de reintento."],
        "rules": ["RN-04", "RN-05", "RN-20"],
        "evidence": "app/dashboard/admin/page.tsx; lib/admin-data.ts; migración 020",
    },
    {
        "module": "Alumnos",
        "id": "CU-05",
        "name": "Consultar el directorio de alumnos",
        "actor": "Administrador",
        "support": "Supabase/PostgreSQL",
        "objective": "Localizar y revisar alumnos con rapidez sin cargar el padrón completo en el navegador.",
        "trigger": "El administrador abre Alumnos o modifica un criterio de consulta.",
        "pre": ["La sesión pertenece a un administrador."],
        "post": ["La tabla presenta únicamente la página y los filtros solicitados."],
        "flow": [
            "La aplicación consulta diez registros mediante paginación real del servidor.",
            "El administrador puede buscar por nombre, apellidos o matrícula.",
            "Puede filtrar por nivel, grado, grupo y estado académico.",
            "Puede ordenar columnas principales en sentido ascendente o descendente.",
            "La tabla muestra identidad institucional, ubicación académica, estado, indicador financiero discreto y acciones.",
            "Los controles Anterior y Siguiente cambian la página sin descargar registros innecesarios.",
        ],
        "alternate": ["No hay coincidencias: se muestra un estado vacío.", "La consulta falla: se informa el error sin ocultar los filtros."],
        "rules": ["RN-04", "RN-06", "RN-07"],
        "evidence": "app/dashboard/admin/alumnos/page.tsx; lib/admin-data.ts",
    },
    {
        "module": "Alumnos",
        "id": "CU-06",
        "name": "Registrar un alumno",
        "actor": "Administrador",
        "support": "Supabase Auth y tabla alumnos",
        "objective": "Crear el expediente académico y la cuenta de acceso de un nuevo alumno.",
        "trigger": "El administrador selecciona Nuevo alumno y confirma el formulario.",
        "pre": ["La matrícula no está asignada a otro alumno.", "El nivel, grado, grupo y sexo son válidos."],
        "post": ["Se crea la cuenta de autenticación y el registro académico.", "Se muestran el correo generado y la contraseña temporal."],
        "flow": [
            "El administrador captura nombre, apellidos, matrícula, nivel, grado o semestre, grupo, sexo y estado.",
            "La aplicación normaliza matrícula y grupo a mayúsculas y valida longitudes y rangos.",
            "El sistema comprueba que la matrícula no exista.",
            "Genera el correo institucional basado en la matrícula y una contraseña temporal.",
            "Crea el usuario aislado en Supabase Auth y después inserta el registro en alumnos.",
            "Actualiza el directorio y muestra las credenciales temporales al administrador.",
        ],
        "alternate": ["Matrícula duplicada o cuenta ya existente: se cancela el alta y se informa la causa.", "Datos académicos fuera de rango: no se envía el formulario."],
        "rules": ["RN-06", "RN-08", "RN-09"],
        "evidence": "components/StudentDirectoryDialogs.tsx; app/dashboard/admin/alumnos/nuevo/page.tsx",
    },
    {
        "module": "Alumnos",
        "id": "CU-07",
        "name": "Consultar el perfil integral de un alumno",
        "actor": "Administrador",
        "support": "Supabase/PostgreSQL",
        "objective": "Concentrar información académica, financiera, de contacto y fiscal de un alumno.",
        "trigger": "El administrador selecciona el icono Ver perfil.",
        "pre": ["El alumno existe y el administrador tiene sesión activa."],
        "post": ["Se presentan los datos disponibles del alumno y del ciclo seleccionado."],
        "flow": [
            "El sistema carga el expediente académico y los pagos del ciclo.",
            "Consulta costos del nivel de cobro y la beca aplicable.",
            "Calcula el concepto activo, el costo efectivo, lo abonado y lo pendiente.",
            "Muestra las mensualidades de agosto a julio con su estado individual.",
            "Muestra tutores, responsables fiscales y el historial paginado de pagos.",
            "Permite navegar ciclos con flechas sin escribir años manualmente.",
        ],
        "alternate": ["Alumno inexistente: se presenta una explicación y un enlace de regreso.", "Ciclo sin costos: se indica que la configuración está pendiente."],
        "rules": ["RN-04", "RN-10", "RN-11", "RN-17"],
        "evidence": "app/dashboard/admin/alumnos/[id]/page.tsx",
    },
    {
        "module": "Alumnos",
        "id": "CU-08",
        "name": "Editar datos y estado académico",
        "actor": "Administrador",
        "support": "Tabla alumnos",
        "objective": "Corregir información académica y administrar la situación institucional del alumno.",
        "trigger": "El administrador abre Editar información o cambia el selector de estado.",
        "pre": ["El alumno existe."],
        "post": ["Los datos válidos quedan actualizados y las vistas dependientes se refrescan."],
        "flow": [
            "El formulario se abre con los valores actuales.",
            "El administrador modifica nombre, apellidos, nivel, grado o semestre y grupo.",
            "El sistema valida campos obligatorios y el máximo de grado del nivel.",
            "El administrador puede elegir activo, pausa temporal o baja definitiva.",
            "La aplicación guarda el cambio y confirma el resultado.",
        ],
        "alternate": ["Rango académico inválido: se rechaza el guardado.", "Estado baja: se detienen nuevos cargos y pagos; el saldo histórico se conserva."],
        "rules": ["RN-08", "RN-12", "RN-13", "RN-14"],
        "evidence": "app/dashboard/admin/alumnos/[id]/page.tsx; components/StudentDirectoryDialogs.tsx; migración 011",
    },
    {
        "module": "Alumnos",
        "id": "CU-09",
        "name": "Administrar tutores de contacto",
        "actor": "Administrador",
        "support": "Función guardar_tutores_alumno",
        "objective": "Registrar los contactos responsables para comunicaciones relacionadas con el alumno.",
        "trigger": "El administrador selecciona Registrar tutor o Editar contactos.",
        "pre": ["El alumno existe."],
        "post": ["Se almacena un contacto principal y, opcionalmente, un segundo contacto."],
        "flow": [
            "La aplicación presenta los contactos actuales, si existen.",
            "El administrador captura relación (madre, padre o tutor), nombre y teléfono del contacto principal.",
            "Puede añadir un segundo contacto con los mismos campos.",
            "El correo es opcional para ambos contactos.",
            "La aplicación valida nombre, teléfono de 10 a 15 dígitos, correo y posiciones únicas.",
            "La función de base de datos reemplaza de forma controlada los contactos del alumno.",
        ],
        "alternate": ["No se proporciona contacto principal: se rechaza la operación.", "Se intentan guardar más de dos contactos o posiciones duplicadas: la base de datos rechaza la solicitud."],
        "rules": ["RN-15", "RN-16"],
        "evidence": "components/GuardianSection.tsx; migración 027",
    },
    {
        "module": "Información fiscal",
        "id": "CU-10",
        "name": "Registrar o editar un responsable fiscal",
        "actor": "Administrador",
        "support": "Catálogos fiscales y función guardar_responsable_fiscal_alumno",
        "objective": "Conservar datos estructurados para emitir CFDI 4.0 en un portal externo.",
        "trigger": "El administrador selecciona Registrar responsable o Editar.",
        "pre": ["Los catálogos de régimen fiscal y uso de CFDI están disponibles.", "El alumno existe."],
        "post": ["El responsable queda relacionado con el alumno y el cambio queda auditado."],
        "flow": [
            "El administrador indica tipo de persona y relación con el alumno.",
            "Captura RFC, código postal fiscal, nombre o razón social, régimen, uso de CFDI y correo opcional.",
            "La interfaz filtra regímenes y usos compatibles con el tipo de persona.",
            "La base de datos normaliza y valida RFC, código postal y catálogos.",
            "El administrador puede señalar un responsable predeterminado.",
            "El sistema guarda la relación y confirma la operación.",
        ],
        "alternate": ["RFC o catálogo incompatible: se rechaza el registro con un mensaje específico.", "Catálogos vacíos o sin permiso: el formulario se deshabilita y explica la migración requerida."],
        "rules": ["RN-17", "RN-18"],
        "evidence": "components/FiscalResponsibleSection.tsx; migraciones 022–024",
    },
    {
        "module": "Información fiscal",
        "id": "CU-11",
        "name": "Copiar información fiscal",
        "actor": "Administrador",
        "support": "Portapapeles del navegador",
        "objective": "Transferir datos fiscales al portal externo sin recaptura manual.",
        "trigger": "El administrador selecciona Copiar junto a un dato o Copiar todo.",
        "pre": ["Existe un responsable fiscal relacionado con el alumno."],
        "post": ["El valor seleccionado queda disponible en el portapapeles y se muestra confirmación visual."],
        "flow": ["El sistema identifica el dato solicitado.", "Escribe el valor en el portapapeles.", "Cambia temporalmente el icono y el texto para confirmar la copia.", "El administrador pega la información en el portal externo de facturación."],
        "alternate": ["El navegador deniega acceso al portapapeles: la aplicación muestra un error y no modifica los datos."],
        "rules": ["RN-19"],
        "evidence": "components/FiscalResponsibleSection.tsx",
    },
    {
        "module": "Configuración financiera",
        "id": "CU-12",
        "name": "Configurar costos y fecha límite de inscripción",
        "actor": "Administrador",
        "support": "Función actualizar_configuracion_escolar",
        "objective": "Definir inscripción, mensualidad y fecha límite por nivel y ciclo.",
        "trigger": "El administrador selecciona un nivel, captura valores y confirma el cambio.",
        "pre": ["El ciclo tiene formato consecutivo AAAA-AAAA.", "No existen pagos para ese nivel y ciclo."],
        "post": ["La configuración se guarda y las deudas del nivel se recalculan."],
        "flow": [
            "El administrador navega al ciclo anterior o siguiente mediante flechas.",
            "Selecciona preescolar, primaria, secundaria o bachillerato.",
            "Captura costos no negativos y una fecha dentro de agosto del año inicial.",
            "El sistema verifica en Supabase si existen pagos del nivel y ciclo.",
            "Muestra un diálogo crítico que advierte el recálculo de deudas.",
            "Tras la confirmación, guarda la configuración, recalcula saldos y aplica la regla de inscripción vencida.",
        ],
        "alternate": ["Existen pagos: los campos quedan deshabilitados y no se permite modificar costos.", "Fecha fuera de agosto o ciclo no consecutivo: se rechaza el guardado."],
        "rules": ["RN-20", "RN-21", "RN-22", "RN-23"],
        "evidence": "app/dashboard/admin/configuracion/page.tsx; migraciones 011–013 y 021",
    },
    {
        "module": "Estado de cuenta",
        "id": "CU-13",
        "name": "Generar y actualizar cargos del ciclo",
        "actor": "Administrador / Sistema",
        "support": "Función generar_estado_cuenta_ciclo y triggers",
        "objective": "Crear cargos individualizados y mantener su monto pagado y estatus sincronizados.",
        "trigger": "Se genera un ciclo o cambia información financiera relevante.",
        "pre": ["Existen costos configurados para los niveles involucrados.", "El ciclo es consecutivo y el día límite es válido."],
        "post": ["Cada cargo tiene concepto, monto esperado, monto pagado, fecha límite y estatus."],
        "flow": [
            "La función valida permisos y parámetros.",
            "Determina los meses agosto–julio y sus años correspondientes.",
            "Genera o actualiza inscripción y mensualidades para alumnos elegibles.",
            "Aplica el costo efectivo de la beca, si corresponde.",
            "Suma los pagos vinculados al concepto, permitiendo pagos acumulados sin perder consistencia.",
            "Clasifica cada cargo como pagado, parcial, vencido o pendiente según saldo y fecha.",
        ],
        "alternate": ["Alumno en pausa o baja: no se generan cargos futuros.", "Configuración faltante: el ciclo no puede calcularse correctamente y se informa la ausencia."],
        "rules": ["RN-10", "RN-12", "RN-24", "RN-25"],
        "evidence": "migraciones 007–009, 011 y 018",
    },
    {
        "module": "Estado de cuenta",
        "id": "CU-14",
        "name": "Aplicar pausa automática por inscripción vencida",
        "actor": "Sistema",
        "support": "Función aplicar_pausas_por_inscripcion_vencida y tarea programada",
        "objective": "Cambiar a pausa a quien no liquide la inscripción antes de la fecha límite.",
        "trigger": "Ejecución diaria programada o actualización del resumen/configuración.",
        "pre": ["Existe fecha límite para el nivel y ciclo actual."],
        "post": ["Los alumnos afectados quedan en pausa automática con fecha de aplicación."],
        "flow": [
            "El sistema obtiene el ciclo vigente y la fecha local de México.",
            "Localiza alumnos activos con inscripción pendiente y fecha límite ya vencida.",
            "Cambia su estado a pausa y registra el origen automático.",
            "Si posteriormente se liquida la inscripción, reactiva automáticamente al alumno.",
            "Una pausa manual o una baja no se revierten por esta regla.",
        ],
        "alternate": ["La extensión de tareas programadas no está disponible: la regla también se ejecuta al consultar el resumen o actualizar configuración."],
        "rules": ["RN-13", "RN-23", "RN-26"],
        "evidence": "migración 021; lib/admin-data.ts",
    },
    {
        "module": "Pagos",
        "id": "CU-15",
        "name": "Registrar un pago desde el perfil",
        "actor": "Administrador",
        "support": "Tabla pagos y triggers financieros",
        "objective": "Registrar un abono válido en el siguiente concepto exigible del alumno.",
        "trigger": "El administrador captura monto y método y selecciona Registrar pago.",
        "pre": ["Existe configuración del nivel y ciclo.", "El alumno puede pagar según su estado.", "Hay un concepto pendiente."],
        "post": ["El pago queda registrado y los saldos, cargos, reportes e historial se actualizan."],
        "flow": [
            "La aplicación determina automáticamente el concepto activo.",
            "Prioriza la inscripción; después recorre mensualidades de agosto a julio.",
            "El administrador captura un monto mayor a cero que no exceda el saldo del concepto.",
            "Selecciona efectivo, tarjeta, transferencia o depósito.",
            "La base de datos valida secuencia, beca, ciclo, nivel de cobro y estado académico.",
            "Se inserta el movimiento y los triggers recalculan deuda y estado de cuenta.",
            "La vista refresca saldos e historial y muestra confirmación.",
        ],
        "alternate": ["El alumno está en pausa automática: solo puede pagar inscripción; al liquidarla se reactiva.", "Pago fuera de secuencia o superior al saldo: la base de datos rechaza la operación.", "Alumno en baja o pausa manual: se bloquea el registro."],
        "rules": ["RN-12", "RN-13", "RN-24", "RN-27", "RN-28"],
        "evidence": "app/dashboard/admin/alumnos/[id]/page.tsx; migraciones 004, 005, 018 y 021",
    },
    {
        "module": "Pagos",
        "id": "CU-16",
        "name": "Registrar cobro rápido",
        "actor": "Administrador",
        "support": "Supabase/PostgreSQL",
        "objective": "Capturar un pago desde el directorio sin abandonar la lista de alumnos.",
        "trigger": "El administrador selecciona el icono de cobro rápido.",
        "pre": ["El alumno está activo y tiene un cargo pendiente."],
        "post": ["El movimiento se registra y el directorio refleja el nuevo estado financiero."],
        "flow": [
            "Se abre un modal sobre la lista.",
            "El sistema consulta saldo, configuración, pagos y beca del ciclo actual en paralelo.",
            "Muestra saldo total, siguiente concepto y saldo máximo permitido.",
            "El administrador captura monto y método de pago.",
            "La aplicación registra el pago y refresca datos del directorio, panel y reportes.",
        ],
        "alternate": ["Sin costos configurados o sin cargos pendientes: el modal informa la situación.", "Alumno no activo: el botón permanece deshabilitado."],
        "rules": ["RN-12", "RN-24", "RN-27", "RN-28"],
        "evidence": "components/StudentDirectoryDialogs.tsx; app/dashboard/admin/alumnos/page.tsx",
    },
    {
        "module": "Pagos",
        "id": "CU-17",
        "name": "Consultar pagos recientes e historial",
        "actor": "Administrador o alumno",
        "support": "Tabla pagos",
        "objective": "Revisar movimientos financieros sin tablas de crecimiento indefinido.",
        "trigger": "Se abre Pagos o el historial del perfil.",
        "pre": ["El usuario tiene permiso de lectura sobre los movimientos solicitados."],
        "post": ["Se presenta una página de movimientos ordenada por fecha descendente."],
        "flow": [
            "La vista de Pagos presenta los movimientos recientes con alumno y datos del cobro.",
            "El historial del perfil consulta diez registros mediante range(from, to).",
            "Muestra fecha, tipo, periodo, método, folio y monto.",
            "El contenedor usa desplazamiento interno y encabezado fijo para facilitar la lectura.",
            "Los controles Anterior y Siguiente solicitan la página correspondiente.",
        ],
        "alternate": ["Sin movimientos: se muestra un estado vacío.", "Error de consulta: se presenta un mensaje sin mezclar datos de otro alumno."],
        "rules": ["RN-04", "RN-29"],
        "evidence": "app/dashboard/admin/pagos/page.tsx; components/PaymentHistory.tsx",
    },
    {
        "module": "Pagos",
        "id": "CU-18",
        "name": "Modificar un pago con auditoría",
        "actor": "Administrador",
        "support": "API privada y función modificar_pago_auditado",
        "objective": "Corregir monto o método de pago preservando trazabilidad e integridad de la secuencia.",
        "trigger": "El administrador abre Modificar pago y confirma los cambios.",
        "pre": ["El pago existe.", "El administrador mantiene una sesión válida."],
        "post": ["El pago se modifica, el cambio se audita y los saldos se recalculan."],
        "flow": [
            "El formulario carga el monto y método actuales.",
            "El administrador modifica al menos uno de ellos.",
            "Captura un motivo de 5 a 500 caracteres y vuelve a escribir su contraseña.",
            "La API valida origen, sesión, rol, contraseña y formato de la solicitud.",
            "La función transaccional valida límites y que no se rompa la secuencia de pagos.",
            "Guarda valores anteriores y nuevos, motivo, usuario y fecha en auditoría.",
            "Recalcula saldos y actualiza la pantalla.",
        ],
        "alternate": ["Contraseña incorrecta: se rechaza la modificación.", "Dejar incompleta inscripción o mensualidad con pagos posteriores: se bloquea el cambio.", "Nuevo monto excede el costo efectivo: se rechaza."],
        "rules": ["RN-01", "RN-28", "RN-30", "RN-31"],
        "evidence": "components/PaymentHistory.tsx; app/api/admin/payments/update/route.ts; migraciones 014 y 016",
    },
    {
        "module": "Pagos",
        "id": "CU-19",
        "name": "Eliminar un pago con auditoría",
        "actor": "Administrador",
        "support": "API privada y función eliminar_pago_auditado",
        "objective": "Retirar un movimiento erróneo sin perder evidencia de la operación eliminada.",
        "trigger": "El administrador selecciona Eliminar pago y confirma el diálogo crítico.",
        "pre": ["El pago existe.", "Se proporcionan motivo y contraseña válidos."],
        "post": ["El pago desaparece de operaciones activas, queda registrado en auditoría y los saldos se recalculan."],
        "flow": [
            "El administrador abre el editor del pago.",
            "Captura el motivo y la contraseña.",
            "La interfaz solicita una segunda confirmación explícita.",
            "La API valida origen, sesión, rol y contraseña.",
            "La función comprueba que eliminar el movimiento no rompa la secuencia ni una promoción.",
            "Copia todos los datos a la bitácora de pagos eliminados y elimina el movimiento.",
            "El sistema recalcula saldos y confirma el resultado.",
        ],
        "alternate": ["Existen mensualidades posteriores: no se elimina la inscripción previa.", "Existen meses posteriores: no se elimina una mensualidad anterior.", "El último pago de inscripción generó una promoción: no se elimina hasta resolver la dependencia."],
        "rules": ["RN-01", "RN-30", "RN-31", "RN-32"],
        "evidence": "components/PaymentHistory.tsx; app/api/admin/payments/update/route.ts; migración 019",
    },
    {
        "module": "Becas",
        "id": "CU-20",
        "name": "Administrar el catálogo de becas",
        "actor": "Administrador",
        "support": "Tabla becas",
        "objective": "Crear, modificar, desactivar o eliminar tipos de beca reutilizables.",
        "trigger": "El administrador abre Configuración de becas y selecciona una acción.",
        "pre": ["La sesión pertenece a un administrador."],
        "post": ["El catálogo refleja el nombre, porcentaje, alcance, descripción y disponibilidad definidos."],
        "flow": [
            "El administrador crea o edita una beca.",
            "Captura nombre, porcentaje entre 0.01 y 100, alcance y descripción opcional.",
            "Indica si está disponible para nuevas asignaciones.",
            "El sistema guarda y actualiza la lista.",
            "Una beca sin dependencias puede eliminarse; una asignada debe conservarse como inactiva.",
        ],
        "alternate": ["Porcentaje inválido: se rechaza el formulario.", "La beca tiene alumnos asignados: la eliminación falla y se indica usar el estado inactivo."],
        "rules": ["RN-33", "RN-34"],
        "evidence": "app/dashboard/admin/becas/configuracion/page.tsx; migración 017",
    },
    {
        "module": "Becas",
        "id": "CU-21",
        "name": "Asignar una beca a un alumno",
        "actor": "Administrador",
        "support": "Función asignar_beca_alumno",
        "objective": "Aplicar un descuento controlado a inscripción, mensualidades o ambos durante un ciclo.",
        "trigger": "El administrador selecciona Asignar beca y confirma el formulario.",
        "pre": ["La beca está activa.", "El alumno no tiene otra beca en el ciclo.", "No existen pagos del alumno en ese ciclo.", "Los costos están configurados."],
        "post": ["La asignación guarda una fotografía del porcentaje y alcance y recalcula cargos y deuda."],
        "flow": [
            "El administrador navega al ciclo correspondiente.",
            "Busca y selecciona un alumno disponible.",
            "Selecciona una beca activa y captura observaciones opcionales.",
            "La función verifica que no existan pagos ni asignación duplicada.",
            "Guarda porcentaje y alcance aplicados para preservar el histórico.",
            "Recalcula inscripción, mensualidades y estado de cuenta del ciclo.",
        ],
        "alternate": ["Ya existen pagos: se rechaza la asignación para evitar reinterpretar movimientos históricos.", "No hay costos configurados: se solicita configurar el nivel primero."],
        "rules": ["RN-24", "RN-33", "RN-35", "RN-36"],
        "evidence": "app/dashboard/admin/becas/page.tsx; migración 018",
    },
    {
        "module": "Becas",
        "id": "CU-22",
        "name": "Retirar una beca",
        "actor": "Administrador",
        "support": "Función retirar_beca_alumno",
        "objective": "Eliminar una asignación antes de que existan pagos dependientes de ella.",
        "trigger": "El administrador selecciona Retirar beca y confirma.",
        "pre": ["La asignación existe.", "No existen pagos del alumno en el ciclo afectado."],
        "post": ["La asignación se elimina y los cargos vuelven al costo sin descuento."],
        "flow": ["La aplicación muestra alumno, beca y ciclo que se modificarán.", "El administrador confirma la acción.", "La función verifica que el ciclo no tenga pagos.", "Elimina la asignación y recalcula deuda y cargos.", "La tabla se refresca y muestra confirmación."],
        "alternate": ["Hay pagos registrados: no se retira la beca y se explica que deben corregirse o eliminarse primero, respetando sus reglas de auditoría y secuencia."],
        "rules": ["RN-30", "RN-35", "RN-36"],
        "evidence": "app/dashboard/admin/becas/page.tsx; migraciones 018 y 019",
    },
    {
        "module": "Reportes",
        "id": "CU-23",
        "name": "Consultar reportes financieros",
        "actor": "Administrador",
        "support": "Estado de cuenta y función de KPIs",
        "objective": "Analizar recaudación, saldo vencido, proyección y alumnos con adeudo.",
        "trigger": "El administrador abre Reportes o cambia pestaña, filtro o página.",
        "pre": ["El estado de cuenta está generado y actualizado."],
        "post": ["La pantalla muestra KPIs y una página de cargos conforme a los criterios."],
        "flow": [
            "El sistema actualiza estatus de cargos y consulta KPIs.",
            "Muestra total recaudado, saldo actual vencido, proyección de ingresos y alumnos con adeudo.",
            "El administrador alterna entre Resumen general, Por cobrar y Estado de cuenta.",
            "Aplica filtros por nivel, grado, grupo y tipo de pago.",
            "La tabla presenta alumno, ubicación académica, concepto, estatus, fecha límite y saldo vencido.",
            "La consulta usa paginación del servidor de diez registros.",
        ],
        "alternate": ["Sin resultados: se muestra una tabla vacía coherente con los filtros.", "Error al actualizar estatus o consultar datos: se muestra un mensaje y los KPIs no se presentan como datos válidos."],
        "rules": ["RN-05", "RN-12", "RN-25", "RN-29", "RN-37"],
        "evidence": "components/FinancialReports.tsx; lib/admin-data.ts; migraciones 007, 009 y 011",
    },
    {
        "module": "Reportes",
        "id": "CU-24",
        "name": "Exportar el reporte a CSV",
        "actor": "Administrador",
        "support": "Navegador",
        "objective": "Obtener una copia tabular de los cargos visibles para análisis o respaldo operativo.",
        "trigger": "El administrador selecciona Exportar a CSV.",
        "pre": ["La consulta del reporte terminó correctamente."],
        "post": ["El navegador descarga un archivo CSV compatible con herramientas de hoja de cálculo."],
        "flow": ["La aplicación construye los encabezados del reporte.", "Convierte cada fila disponible y escapa comas, comillas y saltos.", "Marca el saldo vencido conforme a estatus y fecha.", "Genera el archivo y activa la descarga local."],
        "alternate": ["No hay filas: se conserva la estructura de encabezados o se informa que no existen datos exportables, según el estado de la vista."],
        "rules": ["RN-37"],
        "evidence": "components/FinancialReports.tsx",
    },
    {
        "module": "Reportes",
        "id": "CU-25",
        "name": "Preparar recordatorio por WhatsApp",
        "actor": "Administrador",
        "support": "Tutor principal y WhatsApp",
        "objective": "Abrir un chat dirigido al tutor principal con un recordatorio de adeudo prellenado.",
        "trigger": "El administrador selecciona la campana en un cargo vencido o parcial.",
        "pre": ["El alumno tiene un tutor principal con teléfono válido.", "La fila tiene saldo pendiente."],
        "post": ["WhatsApp se abre en una pestaña nueva con destinatario y mensaje preparados."],
        "flow": [
            "La aplicación consulta el contacto de posición 1 del alumno.",
            "Normaliza el número: conserva código internacional o agrega 52 para diez dígitos de México.",
            "Construye un mensaje con tutor, alumno, concepto, saldo y fecha límite.",
            "Abre https://wa.me/{número}?text={mensaje}.",
            "El administrador revisa el contenido y pulsa Enviar dentro de WhatsApp.",
        ],
        "alternate": ["No existe tutor principal: se solicita registrarlo desde el perfil.", "Teléfono inválido: se informa que debe corregirse.", "El navegador bloquea la ventana: se notifica el bloqueo."],
        "rules": ["RN-15", "RN-38"],
        "evidence": "components/FinancialReports.tsx; migración 027",
    },
    {
        "module": "Portal del alumno",
        "id": "CU-26",
        "name": "Consultar el estado de cuenta personal",
        "actor": "Alumno",
        "support": "Supabase/PostgreSQL",
        "objective": "Permitir que el alumno comprenda si está al corriente y el origen de cualquier saldo vencido.",
        "trigger": "El alumno inicia sesión y abre su panel.",
        "pre": ["La cuenta está asociada a un registro de alumno y no se encuentra en baja."],
        "post": ["Se muestran únicamente los datos financieros autorizados del alumno autenticado."],
        "flow": [
            "El sistema localiza el alumno mediante usuario_id.",
            "Consulta pagos, costos, beca y estado de cuenta del ciclo vigente.",
            "Calcula el saldo vencido como mensualidades vencidas más inscripción pendiente.",
            "Muestra ¡Estás al corriente! si el resultado es cero; de lo contrario, muestra el adeudo vencido.",
            "Resalta el mes actual y distingue mensualidades pagadas, pendientes y futuras.",
            "Muestra inscripción, ciclo, beca y el historial paginado de pagos.",
        ],
        "alternate": ["Cuenta sin alumno asociado: se pide contactar a administración.", "Costos no configurados: se indica que el estado todavía no puede calcularse.", "Error parcial en pagos: el expediente puede mostrarse con una advertencia específica."],
        "rules": ["RN-02", "RN-10", "RN-11", "RN-25", "RN-29"],
        "evidence": "app/dashboard/alumno/page.tsx; components/PaymentHistory.tsx",
    },
    {
        "module": "Portal del alumno",
        "id": "CU-27",
        "name": "Bloquear una cuenta por baja institucional",
        "actor": "Sistema",
        "support": "Layout del portal del alumno",
        "objective": "Impedir que un alumno dado de baja navegue por su portal.",
        "trigger": "Una cuenta de alumno solicita /dashboard/alumno.",
        "pre": ["El usuario está autenticado y su registro académico tiene estado baja."],
        "post": ["El usuario permanece fuera del panel y ve la explicación de suspensión."],
        "flow": ["El layout valida la sesión en el servidor.", "Consulta el estado académico asociado al usuario.", "Detecta el valor baja y redirige a /cuenta-suspendida.", "La pantalla muestra Cuenta suspendida por baja institucional.", "El usuario puede cerrar la sesión y volver al acceso."],
        "alternate": ["Estado activo o pausa: el bloqueo por baja no se aplica; las reglas financieras particulares continúan vigentes."],
        "rules": ["RN-12", "RN-14"],
        "evidence": "app/dashboard/alumno/layout.tsx; app/cuenta-suspendida/page.tsx",
    },
    {
        "module": "Cierre de ciclo",
        "id": "CU-28",
        "name": "Promover automáticamente por reinscripción",
        "actor": "Sistema",
        "support": "Trigger promover_alumno_por_reinscripcion",
        "objective": "Actualizar nivel y grado cuando se paga la inscripción de un ciclo posterior.",
        "trigger": "Se inserta un pago de inscripción para un ciclo escolar nuevo.",
        "pre": ["La promoción está habilitada.", "El ciclo del pago es posterior al ciclo_grado_actual.", "No existe promoción previa para alumno y ciclo."],
        "post": ["El alumno queda en el grado o nivel siguiente y se registra el historial de promoción."],
        "flow": [
            "El trigger determina el ciclo del pago.",
            "Bloquea el registro del alumno para evitar promociones concurrentes.",
            "Verifica que no exista una promoción previa del mismo ciclo.",
            "Incrementa el grado dentro del nivel.",
            "En el máximo de preescolar, primaria o secundaria, cambia al nivel siguiente y reinicia en 1.",
            "Registra nivel y grado anteriores y nuevos en promociones_academicas.",
            "Actualiza ciclo_grado_actual.",
        ],
        "alternate": ["Bachillerato en sexto semestre: se rechaza una promoción adicional.", "Pago del mismo ciclo o anterior: no vuelve a promover.", "promocion_habilitada es falsa: conserva la situación académica."],
        "rules": ["RN-08", "RN-39", "RN-40"],
        "evidence": "lib/academic.ts; migraciones 011 y 026",
    },
]


BUSINESS_RULES = [
    ("RN-01", "Toda operación privada exige una sesión confirmada por Supabase Auth."),
    ("RN-02", "Las rutas administrativas son exclusivas del rol admin; el alumno solo accede a su información autorizada."),
    ("RN-03", "Una sesión no válida, expirada o con rol desconocido no puede continuar en el dashboard."),
    ("RN-04", "Las listas de crecimiento continuo deben consultar páginas limitadas; el tamaño estándar es 10 registros."),
    ("RN-05", "Las métricas solo se consideran válidas después de actualizar los estatus de los cargos."),
    ("RN-06", "La matrícula es única, usa entre 4 y 30 caracteres alfanuméricos y no se modifica desde el editor general."),
    ("RN-07", "La tabla general no muestra cantidades de deuda; solo un indicador financiero discreto."),
    ("RN-08", "Preescolar, primaria y secundaria usan grados; bachillerato usa semestres. Máximos: 3, 6, 3 y 6 respectivamente."),
    ("RN-09", "Al crear un alumno se genera una cuenta de acceso vinculada mediante usuario_id."),
    ("RN-10", "El ciclo escolar comprende agosto del año inicial a julio del año siguiente."),
    ("RN-11", "El saldo vencido del alumno es la suma de mensualidades vencidas más cualquier saldo pendiente de inscripción."),
    ("RN-12", "Los alumnos en baja no generan cargos futuros y no pueden registrar pagos; la deuda histórica se conserva."),
    ("RN-13", "Los alumnos en pausa no generan cargos futuros; una pausa automática por inscripción admite únicamente el pago de inscripción."),
    ("RN-14", "Una cuenta de alumno en baja se redirige a la pantalla de suspensión."),
    ("RN-15", "Debe existir exactamente un contacto principal; puede existir un segundo contacto opcional."),
    ("RN-16", "El teléfono del tutor debe contener de 10 a 15 dígitos y el correo, si se captura, debe ser válido."),
    ("RN-17", "Los datos fiscales se capturan conforme a la constancia y solo el personal administrador puede modificarlos."),
    ("RN-18", "RFC, régimen y uso de CFDI deben ser compatibles con el tipo de persona y los catálogos habilitados."),
    ("RN-19", "La facturación se realiza en un portal externo; el sistema proporciona copia rápida, no emite el CFDI."),
    ("RN-20", "El ciclo escrito en base de datos debe tener formato AAAA-AAAA y representar años consecutivos."),
    ("RN-21", "Los costos no pueden ser negativos."),
    ("RN-22", "Los costos de un nivel y ciclo no se modifican si ya existen pagos registrados para ese nivel de cobro."),
    ("RN-23", "La fecha límite de inscripción debe pertenecer a agosto del año inicial del ciclo."),
    ("RN-24", "La beca reduce el costo efectivo antes de validar pagos y calcular deuda."),
    ("RN-25", "Un cargo es pagado, parcial, vencido o pendiente según monto cubierto y fecha límite."),
    ("RN-26", "Liquidar la inscripción reactiva únicamente una pausa cuyo origen fue el vencimiento de inscripción."),
    ("RN-27", "La inscripción debe liquidarse antes de las mensualidades y estas deben pagarse en orden agosto–julio."),
    ("RN-28", "Ningún pago puede ser menor o igual a cero ni exceder el saldo del concepto."),
    ("RN-29", "Los historiales y tablas extensas deben presentar navegación Anterior/Siguiente y conteo total."),
    ("RN-30", "Modificar o eliminar un pago exige motivo de 5 a 500 caracteres y reautenticación del administrador."),
    ("RN-31", "Toda corrección financiera conserva usuario, fecha, motivo y valores afectados en tablas de auditoría."),
    ("RN-32", "No se elimina un pago si rompe la secuencia de mensualidades, deja inscripción incompleta con pagos posteriores o invalida una promoción."),
    ("RN-33", "Una beca define porcentaje, alcance y disponibilidad; una beca inactiva conserva el histórico."),
    ("RN-34", "Una beca con asignaciones no se elimina físicamente; debe marcarse como inactiva."),
    ("RN-35", "Solo puede existir una beca por alumno y ciclo."),
    ("RN-36", "No se asigna ni retira una beca cuando ya existen pagos del alumno en el ciclo."),
    ("RN-37", "La proyección excluye alumnos dados de baja, mientras su saldo vencido histórico continúa visible."),
    ("RN-38", "El recordatorio dirige el chat al tutor principal y prepara el texto; el administrador realiza el envío final en WhatsApp."),
    ("RN-39", "Una inscripción de un ciclo posterior promueve una sola vez al alumno y registra el cambio."),
    ("RN-40", "Al concluir el máximo de preescolar, primaria o secundaria, la promoción avanza al siguiente nivel; bachillerato termina en sexto semestre."),
]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == 9360, widths
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=None, color=INK, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def build_numbering(doc):
    numbering = doc.part.numbering_part.element

    def create_abstract(abstract_id: int, kind: str):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if kind == "decimal" else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "%1." if kind == "decimal" else "•")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "269")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "271")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Calibri")
        r_fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(r_fonts)
        lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
        abstract.append(lvl)
        numbering.append(abstract)

    create_abstract(910, "decimal")
    create_abstract(911, "bullet")
    return {"decimal": 910, "bullet": 911}


def new_num_id(doc, abstract_id: int):
    numbering = doc.part.numbering_part.element
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(ids or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    n_id = OxmlElement("w:numId")
    n_id.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, n_id])
    p_pr.append(num_pr)


def add_list(doc, items, kind, numbering_ids):
    if not items:
        return
    num_id = new_num_id(doc, numbering_ids[kind])
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, num_id)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    r = p.add_run(value)
    set_run_font(r, size=10.5)
    return p


def add_callout(doc, title, text, color=BLUE, fill=LIGHTER):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(f"{title}. ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_width(cells[idx], widths[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle.paragraph_format.space_after = Pt(8)


def setup_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def setup_headers_footers(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PORTAL ESCOLAR  |  ESPECIFICACIÓN FUNCIONAL")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    first_header = section.first_page_header
    fp = first_header.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [6500, 2860])
    table.style = "Table Grid"
    # Remove visible borders from footer table.
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tbl_borders.append(el)
    table._tbl.tblPr.append(tbl_borders)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r = left.add_run("Casos de uso · Versión 1.0 · Septiembre 2026")
    set_run_font(r, size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.paragraph_format.space_after = Pt(0)
    add_page_field(right)

    first_footer = section.first_page_footer
    fp = first_footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PORTAL ESCOLAR")
    set_run_font(r, size=11, color=SKY, bold=True)
    r.font.all_caps = True

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Documento de Casos de Uso")
    set_run_font(r, size=30, color=NAVY, bold=True)

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema de control escolar y gestión de pagos")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(70)
    r = p.add_run("Especificación funcional para validación operativa")
    set_run_font(r, size=10.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "Alcance del documento",
        "Describe los actores, reglas de negocio, flujos principales, alternativas, excepciones y trazabilidad técnica de las funciones actualmente implementadas en el proyecto.",
        color=BLUE,
        fill=LIGHTER,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Versión 1.0  ·  3 de septiembre de 2026")
    set_run_font(r, size=10.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Preparado para revisión del equipo administrativo y técnico")
    set_run_font(r, size=9.5, color=MUTED)

    doc.add_page_break()


def add_document_control(doc):
    add_heading(doc, "1. Control del documento", 1)
    rows = [
        ("Documento", "Documento de Casos de Uso — Portal Escolar"),
        ("Versión", "1.0"),
        ("Fecha de corte", "3 de septiembre de 2026"),
        ("Estado", "Borrador completo para validación funcional"),
        ("Audiencia", "Dirección, administración escolar, personal de cobranza y equipo técnico"),
        ("Fuente", "Código y migraciones vigentes del repositorio SCP-CEV"),
    ]
    table = add_table(doc, ["Campo", "Valor"], rows, [2700, 6660], 9.5)
    for cell in table.columns[0].cells[1:]:
        set_cell_shading(cell, LIGHTER)
        cell.paragraphs[0].runs[0].font.bold = True

    add_heading(doc, "1.1 Propósito", 2)
    doc.add_paragraph(
        "Este documento define cómo interactúan los usuarios y los procesos automáticos con el Portal Escolar. Su objetivo es servir como base de validación con el equipo, referencia de operación, guía para pruebas de aceptación y punto de partida para futuras mejoras sin confundir funciones existentes con funcionalidades planeadas."
    )

    add_heading(doc, "1.2 Alcance", 2)
    add_list(doc, [
        "Autenticación, protección de rutas, roles y cierre de sesión.",
        "Administración académica de alumnos, estados, tutores y promoción de grado.",
        "Configuración de costos, ciclos, fecha límite de inscripción y pausa automática.",
        "Registro, consulta, modificación y eliminación auditada de pagos.",
        "Catálogo, asignación, aplicación y retiro controlado de becas.",
        "Estado de cuenta, panel administrativo, reportes, exportación y recordatorios.",
        "Consulta del alumno y gestión de responsables fiscales para CFDI 4.0.",
    ], "bullet", doc.numbering_ids)

    add_heading(doc, "1.3 Límites funcionales actuales", 2)
    add_callout(doc, "WhatsApp", "El sistema abre el chat del tutor principal y prepara el mensaje. El administrador debe revisar y pulsar Enviar; no existe envío automatizado por API.", color=AMBER, fill="FFF7ED")
    add_callout(doc, "Facturación", "El sistema almacena y facilita copiar datos fiscales. La emisión del CFDI se realiza en un portal externo y no forma parte de esta aplicación.", color=AMBER, fill="FFF7ED")
    add_callout(doc, "Operación de datos", "Las migraciones 001–027 y la configuración de Supabase deben estar aplicadas. Los datos demo no sustituyen la migración de información real ni las pruebas de aceptación.", color=RED, fill="FEF2F2")


def add_overview(doc):
    add_heading(doc, "2. Visión general del sistema", 1)
    doc.add_paragraph(
        "Portal Escolar es una aplicación web responsiva construida con Next.js, React y Tailwind CSS, con Supabase como plataforma de autenticación y base de datos PostgreSQL. Centraliza el padrón, el estado académico, los cargos, pagos, becas, reportes, tutores y datos fiscales. Las reglas críticas se duplican en la interfaz para mejorar la experiencia y en la base de datos para proteger la integridad aun cuando la solicitud no provenga de la pantalla prevista."
    )

    add_heading(doc, "2.1 Actores", 2)
    rows = [
        ("Administrador", "Gestiona alumnos, costos, pagos, becas, reportes, tutores y datos fiscales."),
        ("Alumno", "Consulta su estado de cuenta, mensualidades, beca e historial autorizado."),
        ("Sistema", "Actualiza estatus, recalcula saldos, aplica pausas y promociones mediante funciones y triggers."),
        ("Tutor principal", "Recibe el recordatorio preparado por el administrador a través de WhatsApp."),
        ("Supabase Auth", "Autentica, mantiene la sesión y permite revalidar la contraseña en acciones financieras críticas."),
        ("WhatsApp", "Servicio externo que abre el chat dirigido y recibe el mensaje prellenado."),
        ("Portal de facturación", "Sistema externo donde el personal pega los datos y emite el CFDI."),
    ]
    add_table(doc, ["Actor", "Responsabilidad"], rows, [2200, 7160], 9.3)

    add_heading(doc, "2.2 Estados y conceptos centrales", 2)
    rows = [
        ("Estado académico", "Activo", "Cursa y puede operar normalmente."),
        ("Estado académico", "Pausa", "Interrupción temporal; no genera cargos futuros. La pausa automática permite liquidar inscripción."),
        ("Estado académico", "Baja", "Baja definitiva; bloquea portal, pagos y cargos futuros, pero conserva históricos."),
        ("Estado de cobro", "Pagado", "Monto pagado igual o superior al esperado."),
        ("Estado de cobro", "Parcial", "Existe abono, pero todavía queda saldo."),
        ("Estado de cobro", "Vencido", "Hay saldo y la fecha límite ya pasó."),
        ("Estado de cobro", "Pendiente", "Hay saldo, pero la fecha límite no ha pasado."),
    ]
    add_table(doc, ["Categoría", "Valor", "Interpretación"], rows, [1850, 1600, 5910], 9.2)

    add_heading(doc, "2.3 Supuestos operativos", 2)
    add_list(doc, [
        "La zona horaria de negocio es America/Mexico_City.",
        "Los ciclos se nombran con años consecutivos, por ejemplo 2026-2027.",
        "La moneda utilizada en la interfaz es peso mexicano (MXN).",
        "La inscripción corresponde a agosto y las mensualidades se recorren de agosto a julio.",
        "El administrador revisa los mensajes y resultados antes de confirmar acciones sensibles.",
    ], "bullet", doc.numbering_ids)


def add_catalog(doc):
    add_heading(doc, "3. Catálogo de casos de uso", 1)
    doc.add_paragraph("La siguiente matriz resume la cobertura. El detalle completo aparece en la sección 4.")
    rows = [(uc["id"], uc["name"], uc["actor"], uc["module"]) for uc in USE_CASES]
    add_table(doc, ["ID", "Caso de uso", "Actor principal", "Módulo"], rows, [900, 3650, 2350, 2460], 8.5)


def add_use_cases(doc):
    add_heading(doc, "4. Especificación detallada", 1)
    doc.add_paragraph(
        "Cada caso usa la misma estructura: objetivo, actores, disparador, condiciones, flujo normal, alternativas y reglas relacionadas. Las rutas y objetos indicados en Evidencia técnica permiten localizar la implementación."
    )
    current_module = None
    for index, uc in enumerate(USE_CASES):
        if uc["module"] != current_module:
            if current_module is not None:
                doc.add_page_break()
            current_module = uc["module"]
            add_heading(doc, current_module, 2)

        add_heading(doc, f'{uc["id"]} — {uc["name"]}', 3)
        add_label_value(doc, "Objetivo", uc["objective"])
        add_label_value(doc, "Actor principal", uc["actor"])
        add_label_value(doc, "Actores o sistemas de apoyo", uc["support"])
        add_label_value(doc, "Disparador", uc["trigger"])

        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run("Precondiciones")
        set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
        add_list(doc, uc["pre"], "bullet", doc.numbering_ids)

        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run("Flujo principal")
        set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
        add_list(doc, uc["flow"], "decimal", doc.numbering_ids)

        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run("Flujos alternos y excepciones")
        set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
        add_list(doc, uc["alternate"], "bullet", doc.numbering_ids)

        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run("Postcondiciones")
        set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
        add_list(doc, uc["post"], "bullet", doc.numbering_ids)

        add_label_value(doc, "Reglas relacionadas", ", ".join(uc["rules"]))
        p = add_label_value(doc, "Evidencia técnica", uc["evidence"])
        p.paragraph_format.space_after = Pt(12)


def add_business_rules(doc):
    doc.add_page_break()
    add_heading(doc, "5. Catálogo de reglas de negocio", 1)
    doc.add_paragraph(
        "Estas reglas son transversales. Cuando una validación aparece tanto en React como en PostgreSQL, la base de datos se considera la última barrera de integridad."
    )
    rows = [(key, text) for key, text in BUSINESS_RULES]
    add_table(doc, ["Regla", "Definición"], rows, [1150, 8210], 8.8)


def add_security_matrix(doc):
    add_heading(doc, "6. Matriz de permisos y controles", 1)
    rows = [
        ("Iniciar/cerrar sesión", "Sí", "Sí", "Supabase Auth; claims; cookies"),
        ("Consultar padrón completo", "Sí", "No", "Middleware y RLS"),
        ("Consultar expediente propio", "Sí", "Sí, solo propio", "RLS por usuario_id"),
        ("Crear/editar alumnos", "Sí", "No", "Rol y políticas de tabla"),
        ("Registrar pagos", "Sí", "No", "RLS, triggers y estado académico"),
        ("Modificar/eliminar pagos", "Sí", "No", "API privada, contraseña, RPC y auditoría"),
        ("Configurar costos", "Sí", "No", "RPC; bloqueo por pagos"),
        ("Administrar becas", "Sí", "No", "RLS y RPC"),
        ("Consultar reportes", "Sí", "No", "RPC protegidas y RLS"),
        ("Gestionar tutores", "Sí", "No", "RPC de administrador; lectura autorizada"),
        ("Gestionar datos fiscales", "Sí", "No", "RPC, catálogos, validación y auditoría"),
        ("Consultar pagos propios", "Sí", "Sí", "RLS por relación con alumnos"),
    ]
    add_table(doc, ["Capacidad", "Administrador", "Alumno", "Control principal"], rows, [2500, 1500, 1500, 3860], 8.8)

    add_heading(doc, "6.1 Controles de integridad financiera", 2)
    add_list(doc, [
        "Secuencia obligatoria: inscripción antes de mensualidades y meses en orden.",
        "Monto positivo y limitado al saldo efectivo después de beca.",
        "Nivel de cobro y ciclo preservados en el movimiento para bloquear correctamente cambios de costos.",
        "Reautenticación, motivo, confirmación visual y auditoría en correcciones y eliminaciones.",
        "Restricciones sobre pagos anteriores cuando existen movimientos posteriores o promociones dependientes.",
        "Recalculo transaccional de deuda, estado de cuenta y estatus después de cambios financieros.",
    ], "bullet", doc.numbering_ids)


def add_traceability(doc):
    add_heading(doc, "7. Matriz de trazabilidad técnica", 1)
    rows = [
        ("CU-01–CU-03", "/login; middleware; Sidebar", "Supabase Auth; claims"),
        ("CU-04", "/dashboard/admin", "obtener_resumen_administrativo; aplicar_pausas_por_inscripcion_vencida"),
        ("CU-05–CU-08", "/dashboard/admin/alumnos; /alumnos/[id]", "alumnos; obtener_filtros_directorio_alumnos"),
        ("CU-09", "GuardianSection", "tutores_alumnos; guardar_tutores_alumno"),
        ("CU-10–CU-11", "FiscalResponsibleSection", "responsables_fiscales; alumnos_responsables_fiscales; catálogos; auditoría"),
        ("CU-12", "/dashboard/admin/configuracion", "configuracion_costos; existen_pagos_nivel_ciclo; actualizar_configuracion_escolar"),
        ("CU-13–CU-14", "Procesos financieros", "estado_cuenta; generar_estado_cuenta_ciclo; actualizar_estatus_estado_cuenta"),
        ("CU-15–CU-17", "Perfil; QuickPaymentModal; PaymentHistory", "pagos; triggers de saldo y estado de cuenta"),
        ("CU-18–CU-19", "/api/admin/payments/update", "modificar_pago_auditado; eliminar_pago_auditado; tablas de auditoría"),
        ("CU-20–CU-22", "/dashboard/admin/becas", "becas; alumnos_becas; asignar_beca_alumno; retirar_beca_alumno"),
        ("CU-23–CU-25", "FinancialReports", "estado_cuenta; KPIs; tutores_alumnos"),
        ("CU-26–CU-27", "/dashboard/alumno; /cuenta-suspendida", "alumnos; pagos; estado_cuenta; configuración; becas"),
        ("CU-28", "Trigger de reinscripción", "promover_alumno_por_reinscripcion; promociones_academicas"),
    ]
    add_table(doc, ["Casos", "Interfaz / ruta", "Datos y servicios"], rows, [1500, 3300, 4560], 8.7)

    add_heading(doc, "7.1 Migraciones por dominio", 2)
    rows = [
        ("001–006", "Base de alumnos, pagos, matrícula, ciclos, costos, secuencia e índices."),
        ("007–010", "Estado de cuenta financiero, sobrepagos controlados, estatus y estado académico."),
        ("011–013", "Reglas de producción, promoción y bloqueo sólido de costos por nivel de cobro."),
        ("014–016", "Edición auditada y método de pago."),
        ("017–019", "Becas aplicadas financieramente y eliminación auditada de pagos."),
        ("020–021", "Dashboard integral, fecha límite y pausa automática por inscripción."),
        ("022–024", "Expediente fiscal CFDI 4.0, permisos de catálogos y validación de RFC."),
        ("025–026", "Preescolar y rangos/promoción de grados y semestres."),
        ("027", "Tutores del alumno y contacto principal para recordatorios."),
    ]
    add_table(doc, ["Migraciones", "Responsabilidad"], rows, [1900, 7460], 9.1)


def add_acceptance(doc):
    add_heading(doc, "8. Guía de validación y aceptación", 1)
    doc.add_paragraph(
        "Antes de reemplazar los datos ficticios por información real, se recomienda ejecutar las siguientes pruebas con un ciclo controlado y conservar evidencia de cada resultado."
    )
    checks = [
        "Un administrador inicia sesión, navega todos los módulos y cierra sesión; un alumno no puede abrir rutas administrativas.",
        "Se registra un alumno de cada nivel y se verifican rangos de grado/semestre, matrícula única y credenciales generadas.",
        "Se capturan uno y dos tutores; se rechazan teléfono, correo o contacto principal inválidos.",
        "Se configuran costos y fecha límite de cada nivel; después de registrar un pago, los campos quedan bloqueados.",
        "Se registra inscripción parcial y total; se comprueba que no se permita mensualidad antes de liquidarla.",
        "Se recorren mensualidades en orden y se intenta un monto superior al saldo para comprobar el rechazo.",
        "Se modifica monto y método con motivo/contraseña; se verifica auditoría y recálculo.",
        "Se elimina un pago permitido y se intenta eliminar otro con dependencias posteriores.",
        "Se crea una beca, se asigna antes de pagar y se comprueba el descuento en cargos y perfil.",
        "Se intenta retirar la beca después de un pago y se valida el mensaje de dependencia.",
        "Se vence una inscripción de prueba y se comprueba pausa automática; al liquidarla, se comprueba reactivación.",
        "Se cambia un alumno a baja y se comprueba bloqueo del portal, pagos y cargos futuros, conservando su saldo histórico.",
        "Se valida que los KPIs, filtros, paginación y CSV coincidan con consultas directas de control.",
        "Se registra tutor principal con número internacional y nacional; ambos abren el chat correcto de WhatsApp.",
        "Se registra persona física y moral con catálogos válidos y se comprueban botones de copia.",
        "Se paga inscripción de un ciclo posterior y se verifica una sola promoción, incluido el cambio de nivel en grado máximo.",
        "Se revisa en teléfono, tableta y escritorio que menús, tarjetas, filtros, tablas y modales sean utilizables.",
    ]
    add_list(doc, checks, "bullet", doc.numbering_ids)

    add_heading(doc, "8.1 Criterio de aprobación", 2)
    add_callout(
        doc,
        "Aprobación funcional",
        "El equipo puede considerar validado un caso cuando su flujo principal termina con la postcondición esperada, sus alternativas críticas muestran mensajes comprensibles y la consulta de control confirma que no existe inconsistencia financiera o acceso no autorizado.",
        color=GREEN,
        fill="F0FDF4",
    )

    add_heading(doc, "8.2 Riesgos y pendientes de operación", 2)
    add_list(doc, [
        "Definir un procedimiento formal de respaldo, restauración y eliminación de datos demo antes de migrar datos reales.",
        "Revisar usuarios administrativos, contraseñas temporales y políticas de cambio de contraseña antes de producción.",
        "Confirmar con el responsable fiscal que los catálogos cargados corresponden a la versión vigente del SAT.",
        "Verificar que la tarea programada de pausa automática esté activa; mantener la ejecución alternativa desde el dashboard como respaldo.",
        "Documentar quién autoriza cambios de costos, bajas, becas y correcciones de pagos.",
        "Definir retención y acceso a bitácoras de pagos y datos fiscales conforme a las políticas de la institución.",
    ], "bullet", doc.numbering_ids)


def add_glossary(doc):
    add_heading(doc, "9. Glosario", 1)
    rows = [
        ("Abono", "Pago parcial aplicado al saldo de un concepto."),
        ("Beca", "Descuento porcentual aplicado a inscripción, mensualidades o ambos dentro de un ciclo."),
        ("Ciclo escolar", "Periodo agosto–julio identificado por dos años consecutivos."),
        ("CFDI", "Comprobante Fiscal Digital por Internet emitido en un portal externo."),
        ("Cargo", "Obligación individual con concepto, monto esperado, monto pagado, fecha límite y estatus."),
        ("Matrícula", "Identificador institucional único del alumno."),
        ("Nivel de cobro", "Nivel académico congelado en el pago para aplicar costos y bloqueos correctos."),
        ("Pausa automática", "Estado temporal originado exclusivamente por inscripción vencida y pendiente."),
        ("RLS", "Row Level Security; políticas de PostgreSQL que limitan qué filas puede leer o modificar cada usuario."),
        ("RPC", "Función PostgreSQL invocada de forma remota desde Supabase."),
        ("Tutor principal", "Contacto en posición 1 utilizado como destinatario del recordatorio."),
    ]
    add_table(doc, ["Término", "Definición"], rows, [2200, 7160], 9.2)

    add_heading(doc, "10. Cierre", 1)
    doc.add_paragraph(
        "La versión descrita cubre el núcleo operativo del control escolar y financiero: identifica al alumno, conserva su situación académica, determina cargos y descuentos, controla la secuencia de pagos, protege correcciones sensibles, ofrece visibilidad administrativa y comunica saldos al responsable registrado. La validación del equipo debe concentrarse ahora en reproducir los escenarios de la sección 8 con datos de prueba y autorizar el procedimiento de migración a datos reales."
    )


def add_toc_field(doc):
    add_heading(doc, "Contenido", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualiza este campo en Word para mostrar el índice con números de página."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    add_callout(doc, "Lectura rápida", "Consulta primero el catálogo de la sección 3. Para validar una operación específica, usa el identificador CU correspondiente y revisa después sus reglas RN relacionadas.", color=BLUE, fill=LIGHTER)
    doc.add_page_break()


def audit_document(doc):
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.50
    assert round(section.page_height.inches, 2) == 11.00
    assert all(round(x.inches, 2) == 1.00 for x in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin))
    assert round(section.header_distance.inches, 3) == 0.492
    assert round(section.footer_distance.inches, 3) == 0.492
    assert doc.styles["Normal"].font.name == "Calibri"
    for table in doc.tables:
        width = table._tbl.tblPr.find(qn("w:tblW"))
        assert width is not None and width.get(qn("w:w")) == "9360"


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    setup_section(doc.sections[0])
    configure_styles(doc)
    doc.numbering_ids = build_numbering(doc)
    setup_headers_footers(doc.sections[0])
    core = doc.core_properties
    core.title = "Documento de Casos de Uso — Portal Escolar"
    core.subject = "Especificación funcional del sistema de control escolar y gestión de pagos"
    core.author = "Equipo del Portal Escolar"
    core.keywords = "casos de uso, control escolar, pagos, alumnos, Supabase, Next.js"
    core.comments = "Documento generado a partir de la implementación vigente del repositorio SCP-CEV."

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    add_cover(doc)
    add_toc_field(doc)
    add_document_control(doc)
    add_overview(doc)
    add_catalog(doc)
    add_use_cases(doc)
    add_business_rules(doc)
    add_security_matrix(doc)
    add_traceability(doc)
    add_acceptance(doc)
    add_glossary(doc)
    audit_document(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
